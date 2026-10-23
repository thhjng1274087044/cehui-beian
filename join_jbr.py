#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""从 _reg.xlsx(合同归档台账) join 经办人 到 工程备案_精修.csv，并生成 建设单位项目联系人(姓氏+先生)。"""
import csv, os, collections
from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
REG = os.path.join(HERE, "_reg.xlsx")
CSV = os.path.join(HERE, "工程备案_精修.csv")

# 复姓
SX = ["欧阳","司马","上官","诸葛","东方","独孤","南宫","令狐","皇甫","宇文","慕容","司徒","夏侯","钟离","宇文"]

def to_contact(name):
    name = (name or "").strip()
    if not name:
        return ""
    # 去掉常见的职务后缀，取人名部分再取姓
    core = name.rstrip("总工经理师监长")
    if not core:
        core = name
    sur = core[:2] if core[:2] in SX else core[0]
    return sur + "先生"

# 1) 读 _reg.xlsx 所有 sheet，header 在第2行，按 合同号 找 经办人
wb = load_workbook(REG, read_only=True, data_only=True)
jbr_map = {}
miss = 0
for ws in wb.worksheets:
    rows = ws.iter_rows(min_row=1, values_only=True)
    header = None
    for ri, row in enumerate(rows, 1):
        if ri == 2:  # header row
            header = list(row)
            # 找列索引
            try:
                i_ht = header.index("合同号")
            except ValueError:
                try: i_ht = header.index("合同编号")
                except ValueError: i_ht = None
            try:
                i_jb = header.index("经办人")
            except ValueError:
                i_jb = None
            if i_ht is None or i_jb is None:
                break
            continue
        if ri < 3:
            continue
        if header is None:
            continue
        c_no = row[i_ht] if i_ht < len(row) else None
        jbr = row[i_jb] if i_jb < len(row) else None
        if c_no:
            c_no = str(c_no).strip()
            jbr = str(jbr).strip() if jbr is not None else ""
            if c_no not in jbr_map and jbr:  # 首次出现的有效经办人
                jbr_map[c_no] = jbr
wb.close()
print(f"经办人映射条数: {len(jbr_map)}")

# 2) 读精修CSV，join + 生成联系人
rows = list(csv.DictReader(open(CSV, encoding="utf-8-sig", newline="")))
fns = list(rows[0].keys())
for c in ["经办人", "建设单位项目联系人"]:
    if c not in fns:
        fns.append(c)

matched = 0
for r in rows:
    cid = (r.get("合同号") or "").strip()
    jbr = jbr_map.get(cid, "")
    r["经办人"] = jbr
    r["建设单位项目联系人"] = to_contact(jbr)
    if jbr:
        matched += 1

with open(CSV, "w", encoding="utf-8-sig", newline="") as f:
    w = csv.DictWriter(f, fieldnames=fns)
    w.writeheader(); w.writerows(rows)

print(f"精修表行数: {len(rows)} | 匹配到经办人: {matched} | 未匹配: {len(rows)-matched}")
# 抽样检查
print("\n抽样(前10条匹配):")
n=0
for r in rows:
    if r["经办人"]:
        print(f"  {r['合同号']} | 经办人:{r['经办人']} -> 联系人:{r['建设单位项目联系人']}")
        n+=1
        if n>=10: break
